import os
import sys
import subprocess
from pathlib import Path

# Ensure execution inside .venv if run with an external python interpreter
if not hasattr(sys, "real_prefix") and (sys.base_prefix == sys.prefix):
    venv_python = Path(__file__).resolve().parent.parent / ".venv" / "Scripts" / "python.exe"
    if venv_python.exists() and sys.executable.lower() != str(venv_python).lower():
        result = subprocess.run([str(venv_python), *sys.argv], check=False)
        sys.exit(result.returncode)

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


def build_holdout():
    raw_csv = Path(__file__).resolve().parent.parent / "data" / "raw" / "January2015toNovember2025.csv"
    output_xlsx = Path(__file__).resolve().parent.parent / "data" / "osha_holdout_unlabeled.xlsx"

    print("=" * 70)
    print("Building OSHA SIR Holdout Labeling Workbook")
    print("=" * 70)

    # 1. Load raw CSV
    if not raw_csv.exists():
        print(f"ERROR: Raw data file not found: {raw_csv}", file=sys.stderr)
        sys.exit(1)

    print(f"[1] Loading {raw_csv.name}...")
    df = pd.read_csv(raw_csv, low_memory=False, encoding="utf-8")
    initial_count = len(df)
    print(f"    Initial raw rows: {initial_count:,}")

    # 2. DROP sensitive / PII / location columns entirely
    drop_cols = ["Employer", "Address1", "Address2", "City", "Latitude", "Longitude", "Zip"]
    existing_drop_cols = [c for c in drop_cols if c in df.columns]
    df = df.drop(columns=existing_drop_cols)
    print(f"[2] Dropped columns: {', '.join(existing_drop_cols)}")

    # 3. Filter Primary NAICS starting with 2111, 2131, 486
    # 2111: Oil and Gas Extraction
    # 2131: Support Activities for Mining / Oil and Gas
    # 486:  Pipeline Transportation
    naics_str = df["Primary NAICS"].dropna().astype(str).str.strip()
    naics_mask = df.index.isin(
        naics_str[naics_str.str.startswith(("2111", "2131", "486"))].index
    )
    df_naics = df[naics_mask].copy()
    naics_surviving = len(df_naics)
    print(f"[3] Filtered Primary NAICS (starts with 2111, 2131, 486):")
    print(f"    Surviving rows: {naics_surviving:,} (retained {naics_surviving / initial_count * 100:.2f}%)")

    # 4. Filter Final Narrative length between 250 and 1800 characters
    narr_series = df_naics["Final Narrative"].fillna("").astype(str).str.strip()
    narr_len = narr_series.str.len()
    len_mask = (narr_len >= 250) & (narr_len <= 1800)
    df_len = df_naics[len_mask].copy()
    len_surviving = len(df_len)
    print(f"[4] Filtered Final Narrative length [250, 1800] characters:")
    print(f"    Surviving rows: {len_surviving:,} (retained {len_surviving / naics_surviving * 100:.2f}%)")

    # 5. Create severity_bucket column: "high" if Amputation == 1 or Loss of Eye == 1, else "hospitalized"
    is_high = (df_len["Amputation"] == 1) | (df_len["Loss of Eye"] == 1)
    df_len["severity_bucket"] = "hospitalized"
    df_len.loc[is_high, "severity_bucket"] = "high"

    high_count = (df_len["severity_bucket"] == "high").sum()
    hosp_count = (df_len["severity_bucket"] == "hospitalized").sum()
    print(f"[5] Created severity_bucket:")
    print(f"    - 'high' (Amputation == 1 or Loss of Eye == 1) : {high_count:,}")
    print(f"    - 'hospitalized'                               : {hosp_count:,}")

    # 6. Stratified sampling: 150 rows (60% high = 90 rows, 40% hospitalized = 60 rows)
    target_high = 90
    target_hosp = 60
    if high_count < target_high or hosp_count < target_hosp:
        print(f"WARNING: Insufficient rows in buckets ({high_count} high, {hosp_count} hospitalized).", file=sys.stderr)
        target_high = min(target_high, high_count)
        target_hosp = min(target_hosp, hosp_count)

    sample_high = df_len[df_len["severity_bucket"] == "high"].sample(n=target_high, random_state=42)
    sample_hosp = df_len[df_len["severity_bucket"] == "hospitalized"].sample(n=target_hosp, random_state=42)

    # Combine and shuffle for a balanced distribution
    sample_df = pd.concat([sample_high, sample_hosp]).sample(frac=1.0, random_state=42).reset_index(drop=True)
    print(f"[6] Sampled 150 rows (random_state=42):")
    print(f"    - High severity : {len(sample_high)} ({len(sample_high) / len(sample_df) * 100:.1f}%)")
    print(f"    - Hospitalized  : {len(sample_hosp)} ({len(sample_hosp) / len(sample_df) * 100:.1f}%)")
    print(f"    - Total sample  : {len(sample_df)}")

    # 7. Prepare final columns
    # Columns in exact order:
    # id (1-150), source_id (original ID), EventDate, Primary NAICS, EventTitle, SourceTitle, NatureTitle, Final Narrative,
    # BLANK columns: energy_type, barrier_failure, sif_potential, lsr_tag, labeled_by, notes
    out_df = pd.DataFrame()
    out_df["id"] = range(1, len(sample_df) + 1)
    out_df["source_id"] = sample_df["ID"]
    out_df["EventDate"] = sample_df["EventDate"]
    out_df["Primary NAICS"] = sample_df["Primary NAICS"]
    out_df["EventTitle"] = sample_df["EventTitle"]
    out_df["SourceTitle"] = sample_df["SourceTitle"]
    out_df["NatureTitle"] = sample_df["NatureTitle"]
    out_df["Final Narrative"] = sample_df["Final Narrative"]

    # Blank labeling columns
    blank_cols = ["energy_type", "barrier_failure", "sif_potential", "lsr_tag", "labeled_by", "notes"]
    for c in blank_cols:
        out_df[c] = ""

    # Verify no sensitive columns are present
    for forbidden in drop_cols:
        assert forbidden not in out_df.columns, f"Security check failed: {forbidden} found in output!"

    # 8. Create and style Excel Workbook with openpyxl
    print(f"[7] Writing workbook to {output_xlsx.name} with formatting & validations...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OSHA_Holdout_Unlabeled"

    # Write headers
    headers = list(out_df.columns)
    ws.append(headers)

    # Write data rows
    for row in out_df.itertuples(index=False):
        ws.append(list(row))

    # Freeze top row
    ws.freeze_panes = "A2"

    # Define styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Slate dark blue
    data_font = Font(name="Calibri", size=10)
    border_side = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Apply header formatting
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    ws.row_dimensions[1].height = 28

    # Apply row and cell formatting
    for row_idx in range(2, len(out_df) + 2):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border

            header_name = headers[col_idx - 1]
            if header_name == "Final Narrative":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif header_name in ["id", "source_id", "EventDate", "Primary NAICS", "sif_potential"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Column widths configuration
    col_widths = {
        "id": 8,
        "source_id": 14,
        "EventDate": 13,
        "Primary NAICS": 15,
        "EventTitle": 32,
        "SourceTitle": 28,
        "NatureTitle": 28,
        "Final Narrative": 80,  # Explicitly set to 80
        "energy_type": 16,
        "barrier_failure": 28,
        "sif_potential": 15,
        "lsr_tag": 26,
        "labeled_by": 14,
        "notes": 30,
    }

    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        width = col_widths.get(col_name, 20)
        ws.column_dimensions[col_letter].width = width

    # 9. Add Data Validation Dropdowns
    max_row = len(out_df) + 1

    # 9a. sif_potential dropdown: TRUE, FALSE
    sif_col_letter = get_column_letter(headers.index("sif_potential") + 1)
    dv_sif = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    dv_sif.error = "Please select TRUE or FALSE"
    dv_sif.errorTitle = "Invalid Selection"
    ws.add_data_validation(dv_sif)
    dv_sif.add(f"{sif_col_letter}2:{sif_col_letter}{max_row}")

    # 9b. lsr_tag dropdown: 9 IOGP rules + "None"
    iogp_rules = [
        "Bypassing Safety Controls",
        "Confined Space",
        "Driving",
        "Energy Isolation",
        "Hot Work",
        "Line of Fire",
        "Safe Mechanical Lifting",
        "Work Authorisation",
        "Working at Height",
        "None",
    ]
    lsr_options = ",".join(iogp_rules)
    lsr_col_letter = get_column_letter(headers.index("lsr_tag") + 1)
    dv_lsr = DataValidation(type="list", formula1=f'"{lsr_options}"', allow_blank=True)
    dv_lsr.error = "Please select a valid Life-Saving Rule or None"
    dv_lsr.errorTitle = "Invalid LSR"
    ws.add_data_validation(dv_lsr)
    dv_lsr.add(f"{lsr_col_letter}2:{lsr_col_letter}{max_row}")

    # 9c. energy_type dropdown: electrical, height, mechanical, pressure, thermal, vehicle, toxic, none
    energy_types = "electrical,height,mechanical,pressure,thermal,vehicle,toxic,none"
    energy_col_letter = get_column_letter(headers.index("energy_type") + 1)
    dv_energy = DataValidation(type="list", formula1=f'"{energy_types}"', allow_blank=True)
    dv_energy.error = "Please select a valid energy category"
    dv_energy.errorTitle = "Invalid Energy Type"
    ws.add_data_validation(dv_energy)
    dv_energy.add(f"{energy_col_letter}2:{energy_col_letter}{max_row}")

    # Save workbook
    wb.save(output_xlsx)
    print(f"\n[+] SUCCESS: Saved labeling workbook to {output_xlsx}")
    print(f"    Columns ({len(headers)}): {', '.join(headers)}")
    print(f"    Rows: {len(out_df)}")


if __name__ == "__main__":
    build_holdout()
